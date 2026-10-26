#!/usr/bin/env python3
"""
PAUT 결함 기록 양식(엑셀) 생성기.

입력
  --columns  칸 정의 JSON. app/tests/labels/templateColumns.test.ts 를 PAUT_LABEL_COLUMNS_OUT=<경로> 로 돌리면 생긴다.
             칸 정의의 단일 출처는 app/src/shared/labels/columns.ts 이다 (여기서 칸을 따로 정하지 않는다).
  --prefill  미리 채울 값 JSON (기본: 이 폴더의 prefill_x3.json). 열거 칸은 코드로 적고 여기서 표시 이름으로 바꾼다.
  --docx     (선택) 00_결함정보.docx. 주면 prefill 의 물리결함 11건이 문서 글과 같은지 먼저 확인한다.
  --blind    판독자용(블라인드) 양식: '물리결함' 시트를 비우고, 판독 단계를 '블라인드'로 채운다.
             주지 않으면 대조용 양식: 결함정보 11건을 채우고 판독 단계를 '설계값 대조'로 채운다(블라인드 판독자에게 주지 않음).
출력
  --out      xlsx. 같은 폴더의 임시 파일에 저장해 다시 열어 확인(시트·머리글·드롭다운·입력 검사·조건부 서식·미리 채운 값·
             안내 문구)한 뒤에만 --out 으로 바꿔 넣는다. 틀리면 종료 코드 1 이고 --out 은 그대로다.
             --out 이 이미 있고 사람이 적은 값(미리 채운 값과 다른 칸)이 있으면 덮어쓰지 않는다(--force 로만).

필요: python3, openpyxl 3.1 (새로 설치하는 것 없음). 다시 만드는 명령은 README.md.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

FORM_VERSION = 'v1'
COLUMNS_FORMAT = 'paut-label-columns'
COLUMNS_FORMAT_VERSION = 2
GUIDE_SHEET = '안내'
LIST_SHEET = '목록'
# 드롭다운·형식을 거는 줄 수 (머리글 제외). 이보다 많이 적으면 아래 줄은 드롭다운이 없다 → 안내에 적음
DATA_ROWS = 500
FIRST_DATA_ROW = 2
LAST_DATA_ROW = FIRST_DATA_ROW + DATA_ROWS - 1
COMMENT_AUTHOR = 'PAUT 라벨 양식'
# 이 글자가 안내 시트에 있으면 오래된 양식이다 (지침서가 생기기 전 문구)
STALE_GUIDE_TEXT = '작성 중'

FILL_REQUIRED = PatternFill('solid', fgColor='FCE4D6')  # 주황 = 필수
FILL_APP = PatternFill('solid', fgColor='DDEBF7')  # 파랑 = 비워도 됨(앱이 채움)
FILL_OPTIONAL = PatternFill('solid', fgColor='EDEDED')  # 회색 = 선택
FILL_MISSING = PatternFill('solid', fgColor='FFC7CE', bgColor='FFC7CE')  # 빨강 = 조건부 필수인데 빔
HEADER_BORDER = Border(bottom=Side(style='medium', color='808080'))
PHYSICAL_TAB_COLOR = 'F4B183'
READER_TAB_COLOR = 'A9D08E'

FORMAT_LABEL = {
    'number': '숫자 (단위 없이)',
    'text': '글자',
    'enum': '목록에서 고름',
    'bool': '예 / 아니오',
    'time': '날짜·시각 (예: 2026-10-20 14:05)',
}

BOOL_LIST = 'L_BOOL'
# 다른 시트의 열을 드롭다운으로 (같은 값을 쓰게 돕는 것. 목록에 없어도 경고만 하고 받는다)
REFERENCE_LISTS = [
    # (드롭다운을 걸 시트, 칸 key, 값을 가져올 시트, 칸 key, 이름)
    ('indication', 'fileName', 'file', 'fileName', 'L_FILE_NAMES'),
    ('review', 'fileName', 'file', 'fileName', 'L_FILE_NAMES'),
    ('worktime', 'fileName', 'file', 'fileName', 'L_FILE_NAMES'),
    ('physical', 'specimenId', 'file', 'specimenId', 'L_SPECIMEN_IDS'),
]
# 판독 단계: 판독자용 = 블라인드, 대조용 = 설계값 대조
STAGE_BLIND = 'blind'
STAGE_RECONCILED = 'reconciled'


class TemplateError(Exception):
    pass


# ───────────────────────────── 입력 ─────────────────────────────


def load_columns(path: Path) -> dict:
    data = json.loads(path.read_text(encoding='utf-8'))
    if data.get('format') != COLUMNS_FORMAT:
        raise TemplateError(f"칸 정의 JSON 형식이 아닙니다: {path} (format = {data.get('format')!r})")
    if data.get('formatVersion') != COLUMNS_FORMAT_VERSION:
        raise TemplateError(
            f"칸 정의 JSON 버전 {data.get('formatVersion')} 은 이 생성기({COLUMNS_FORMAT_VERSION})와 다릅니다. "
            '테스트(templateColumns.test.ts)로 JSON 을 다시 만드십시오.'
        )
    names = [s['name'] for s in data['sheets']]
    for reserved in (GUIDE_SHEET, LIST_SHEET):
        if reserved in names:
            raise TemplateError(f"데이터 시트 이름 '{reserved}' 는 생성기가 쓰는 이름입니다.")
    return data


def columns_fingerprint(columns: dict) -> str:
    canon = json.dumps(columns, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
    return hashlib.sha256(canon.encode('utf-8')).hexdigest()[:12]


def sheet_by_id(columns: dict, sheet_id: str) -> dict:
    for s in columns['sheets']:
        if s['id'] == sheet_id:
            return s
    raise TemplateError(f"칸 정의에 시트 '{sheet_id}' 가 없습니다.")


def column_index(sheet: dict, key: str) -> int:
    """1부터 세는 열 번호"""
    for i, c in enumerate(sheet['columns'], start=1):
        if c['key'] == key:
            return i
    raise TemplateError(f"'{sheet['name']}' 시트에 칸 '{key}' 가 없습니다 (columns.ts 와 생성기를 함께 고쳐야 합니다).")


def column_by_key(sheet: dict, key: str) -> dict:
    return sheet['columns'][column_index(sheet, key) - 1]


def last_row(sheet: dict) -> int:
    """드롭다운·검사를 거는 마지막 줄 (판독자 시트는 한 줄)"""
    n = sheet.get('maxRows')
    return FIRST_DATA_ROW + (n if n else DATA_ROWS) - 1


def data_range(sheet: dict, col_no: int) -> str:
    """그 열의 데이터 칸 범위 ('E2:E501'). 한 줄이면 'E2' (엑셀·openpyxl 이 그렇게 줄여 저장함)"""
    letter = get_column_letter(col_no)
    last = last_row(sheet)
    return f'{letter}{FIRST_DATA_ROW}' if last == FIRST_DATA_ROW else f'{letter}{FIRST_DATA_ROW}:{letter}{last}'


def enum_label(col: dict, code: object) -> str:
    for v in col['enumValues']:
        if v['code'] == code:
            return v['label']
    raise TemplateError(f"칸 '{col['id']}' 에 코드 {code!r} 가 없습니다. 고를 수 있는 코드: {[v['code'] for v in col['enumValues']]}")


def cell_value(col: dict, value: object, bool_labels: dict) -> object:
    """미리 채울 값 → 엑셀 칸 값 (열거 = 표시 이름, 참거짓 = 예/아니오)"""
    fmt = col['format']
    if fmt == 'enum':
        return enum_label(col, value)
    if fmt == 'bool':
        if not isinstance(value, bool):
            raise TemplateError(f"칸 '{col['id']}' 는 참/거짓입니다 (받은 값 {value!r}).")
        return bool_labels['true' if value else 'false']
    if fmt == 'number':
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise TemplateError(f"칸 '{col['id']}' 는 숫자입니다 (받은 값 {value!r}).")
        return value
    return str(value)


def prefill_cells(columns: dict, prefill: dict, stage: str | None = None) -> dict[str, list[dict[int, object]]]:
    """시트 id → 줄마다 {열 번호: 엑셀 값}. 칸 key·필수 칸·코드를 확인한다. stage 를 주면 판독자 시트의 판독 단계를 채운다"""
    out: dict[str, list[dict[int, object]]] = {}
    for sheet_id, rows in prefill.get('sheets', {}).items():
        sheet = sheet_by_id(columns, sheet_id)
        by_key = {c['key']: c for c in sheet['columns']}
        if sheet.get('maxRows') and len(rows) > sheet['maxRows']:
            raise TemplateError(f"미리 채울 값 '{sheet['name']}' 는 {sheet['maxRows']}줄까지입니다 (받은 {len(rows)}줄).")
        lines: list[dict[int, object]] = []
        for n, row in enumerate(rows, start=1):
            row = dict(row)
            if sheet_id == 'reader' and stage is not None:
                row['readingStage'] = stage
            cells: dict[int, object] = {}
            for key, value in row.items():
                if key.startswith('_'):
                    continue
                if key not in by_key:
                    raise TemplateError(f"미리 채울 값 '{sheet['name']}' {n}번째 줄: 칸 '{key}' 가 칸 정의에 없습니다.")
                cells[column_index(sheet, key)] = cell_value(by_key[key], value, columns['boolLabels'])
            if not row.get('_partial'):
                for c in sheet['columns']:
                    if c['required'] and column_index(sheet, c['key']) not in cells:
                        raise TemplateError(f"미리 채울 값 '{sheet['name']}' {n}번째 줄: 필수 칸 '{c['header']}' 가 비었습니다.")
            lines.append(cells)
        out[sheet_id] = lines
    return out


def mode_rows(columns: dict, prefill: dict, blind: bool) -> dict[str, list[dict[int, object]]]:
    rows = prefill_cells(columns, prefill, STAGE_BLIND if blind else STAGE_RECONCILED)
    if blind:
        rows.pop('physical', None)
    return rows


# ───────────────────────────── 결함정보.docx 대조 ─────────────────────────────

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
_DEFECT_LINE = re.compile(r'^(?P<text>.+?)\s*[-–—]\s*(?P<pos>\d+(?:\.\d+)?)\s*(?:mm)?\s*$', re.IGNORECASE)


def docx_lines(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read('word/document.xml'))
    lines = []
    for p in root.iter(f'{_W}p'):
        t = ''.join(x.text or '' for x in p.iter(f'{_W}t')).strip()
        if t:
            lines.append(t)
    return lines


def docx_defects(path: Path) -> list[tuple[str, str, float]]:
    """(파일 이름 줄, 종류 글, 위치) 목록. 위치 숫자로 끝나지 않는 줄은 파일 이름 줄로 본다"""
    current = None
    out = []
    for line in docx_lines(path):
        m = _DEFECT_LINE.match(line)
        if m and current is not None:
            out.append((current, m.group('text').strip(), float(m.group('pos'))))
        else:
            current = line
    return out


def _norm(s: str) -> str:
    return re.sub(r'\s+', ' ', s).strip().casefold()


def check_docx(path: Path, prefill: dict) -> int:
    doc = [(_norm(f), _norm(t), p) for f, t, p in docx_defects(path)]
    rows = prefill.get('sheets', {}).get('physical', [])
    mine = []
    for r in rows:
        d = r.get('_docx')
        if d is None:
            raise TemplateError("물리결함 미리 채울 값에 '_docx' 출처가 없는 줄이 있습니다.")
        if r.get('centerMm') != d['positionMm']:
            raise TemplateError(f"{d['file']} {d['text']}: 중심 {r.get('centerMm')} ≠ 문서 위치 {d['positionMm']}")
        mine.append((_norm(d['file']), _norm(d['text']), float(d['positionMm'])))
    if doc != mine:
        diff = [f'  문서 {a}  /  prefill {b}' for a, b in zip(doc, mine) if a != b]
        raise TemplateError(
            f'결함정보.docx 와 prefill 이 다릅니다 (문서 {len(doc)}건, prefill {len(mine)}건).\n' + '\n'.join(diff)
        )
    return len(doc)


# ───────────────────────────── 입력 검사 (칸마다 하나) ─────────────────────────────


def _num(v: float) -> str:
    return str(int(v)) if float(v).is_integer() else repr(float(v))


def rule_validation(columns: dict, sheet: dict, col: dict) -> tuple[str, str, str | None, str] | None:
    """숫자·시각 칸의 데이터 유효성 (형식, formula1, formula2, 오류 문구). 걸 것이 없으면 None.
    custom 수식은 그 열 2행 기준 상대 참조다 (엑셀이 아래 줄로 옮겨 적용)."""
    fmt = col['format']
    if fmt not in ('number', 'time'):
        return None
    rule = col.get('rule') or {}
    me = f"{get_column_letter(column_index(sheet, col['key']))}{FIRST_DATA_ROW}"
    parts = [f'ISNUMBER({me})']
    msgs: list[str] = []
    lo, hi = rule.get('min'), rule.get('max')
    if fmt == 'number':
        if rule.get('int'):
            parts.append(f'INT({me})={me}')
            msgs.append('정수')
        if lo is not None:
            parts.append(f'{me}>={_num(lo)}')
        if hi is not None and not rule.get('maxThicknessTimes'):
            parts.append(f'{me}<={_num(hi)}')
        if lo is not None or hi is not None:
            msgs.append(f"{_num(lo) if lo is not None else ''}~{_num(hi) if hi is not None else ''}")
        k = rule.get('maxThicknessTimes')
        if k:
            fs = sheet_by_id(columns, 'file')
            ti = column_index(fs, 'thicknessMm')
            fc = get_column_letter(column_index(sheet, 'fileName'))
            rng = f"'{fs['name']}'!$A${FIRST_DATA_ROW}:${get_column_letter(ti)}${LAST_DATA_ROW}"
            look = f'VLOOKUP(${fc}{FIRST_DATA_ROW},{rng},{ti},FALSE)'
            mult = '' if k == 1 else f'*{_num(k)}'
            parts.append(f'{me}<=IF(IFERROR({look},0)>0,{look}{mult},{_num(hi)})')
            msgs.append('실제 깊이는 그 파일 두께 이하' if k == 1 else f'펼친 깊이는 그 파일 두께의 {_num(k)}배 이하')
    other = rule.get('atLeastKey')
    if other:
        oc = f"{get_column_letter(column_index(sheet, other))}{FIRST_DATA_ROW}"
        parts.append(f'OR(ISBLANK({oc}),{me}>={oc})')
        msgs.append(f"'{column_by_key(sheet, other)['header']}' 이상")
    if fmt == 'number' and not other and not rule.get('maxThicknessTimes'):
        # 단순 범위는 엑셀 기본 숫자 검사로 (사용자에게 범위가 바로 보인다)
        if lo is None or hi is None:
            raise TemplateError(f"숫자 칸 '{col['id']}' 에 범위(min·max)가 없습니다 (columns.ts rule).")
        kind = 'whole' if rule.get('int') else 'decimal'
        msg = f"숫자만, {_num(lo)}~{_num(hi)}{' 정수' if rule.get('int') else ''}. 단위는 머리글에 있으니 빼고 적습니다."
        return (kind, _num(lo), _num(hi), msg)
    formula = 'AND(' + ','.join(parts) + ')'
    if len(formula) > 255:
        raise TemplateError(f"칸 '{col['id']}' 의 검사 수식이 엑셀 한도(255자)를 넘습니다: {len(formula)}자")
    if fmt == 'number':
        msg = f"숫자만{', ' + ', '.join(msgs) if msgs else ''}. 단위는 머리글에 있으니 빼고 적습니다."
    else:
        msg = f"날짜·시각으로 적습니다 (예: 2026-10-20 14:05){', ' + ', '.join(msgs) if msgs else ''}."
    return ('custom', formula, None, msg)


def expected_validations(columns: dict, sheet: dict, col: dict) -> list[tuple[str, str]]:
    """칸에 걸려야 하는 (형식, formula1) 목록 — 만들 때와 확인할 때 같은 함수를 쓴다"""
    want: list[tuple[str, str]] = []
    if col['format'] == 'enum':
        want.append(('list', list_name(col['enumTable'])))
    elif col['format'] == 'bool':
        want.append(('list', BOOL_LIST))
    else:
        rv = rule_validation(columns, sheet, col)
        if rv is not None:
            want.append((rv[0], rv[1]))
    for r in REFERENCE_LISTS:
        if (r[0], r[1]) == (sheet['id'], col['key']):
            want.append(('list', r[4]))
    return want


def required_when_formula(sheet: dict, col: dict) -> str | None:
    """조건부 필수 칸이 조건에서 비었을 때 빨갛게 칠할 수식 (2행 기준 상대 참조)"""
    rw = col.get('requiredWhen')
    if not rw:
        return None
    cond_col = column_by_key(sheet, rw['key'])
    cl = get_column_letter(column_index(sheet, rw['key']))
    me = f"{get_column_letter(column_index(sheet, col['key']))}{FIRST_DATA_ROW}"
    labels = [enum_label(cond_col, code) for code in rw['codes']]
    conds = [f'${cl}{FIRST_DATA_ROW}="{lab}"' for lab in labels]
    cond = conds[0] if len(conds) == 1 else 'OR(' + ','.join(conds) + ')'
    return f'AND({cond},ISBLANK({me}))'


def required_when_text(sheet: dict, col: dict) -> str | None:
    rw = col.get('requiredWhen')
    if not rw:
        return None
    cond_col = column_by_key(sheet, rw['key'])
    labels = ' / '.join(enum_label(cond_col, c) for c in rw['codes'])
    tail = f" (비면 '{enum_label(col, rw['fallback'])}')" if rw.get('fallback') else ''
    return f"'{cond_col['header']}' = {labels} 이면 필수{tail}"


# ───────────────────────────── 양식 만들기 ─────────────────────────────


def text_width(s: str) -> int:
    """엑셀 열 너비 어림 (한글·전각 = 2)"""
    return sum(2 if ord(ch) >= 0x1100 else 1 for ch in s)


def column_width(col: dict) -> float:
    w = text_width(col['header']) + 3
    if col['format'] == 'enum':
        w = max(w, min(max(text_width(v['label']) for v in col['enumValues']) + 4, 30))
    if col['format'] == 'time':
        w = max(w, 18)
    if col['key'] == 'note':
        w = 48
    if col['key'] == 'fileName':
        w = max(w, 24)
    return float(min(max(w, 10), 48))


def header_status(col: dict) -> tuple[str, PatternFill]:
    if col['required']:
        return '[필수]', FILL_REQUIRED
    if col.get('appFilled'):
        return '[비워도 됨 · 앱이 가져올 때 채움]', FILL_APP
    if col.get('requiredWhen'):
        return '[조건부 필수]', FILL_OPTIONAL
    return '[선택]', FILL_OPTIONAL


def empty_value_text(col: dict, bool_labels: dict) -> str | None:
    v = col.get('emptyValue')
    if v is None:
        return None
    if col['format'] == 'enum':
        return enum_label(col, v)
    if col['format'] == 'bool':
        return bool_labels['true' if v else 'false']
    return '빈 글자' if v == '' else str(v)


def header_comment(columns: dict, sheet: dict, col: dict) -> str:
    status, _ = header_status(col)
    lines = [f"{status} {col['help']}", f"형식: {FORMAT_LABEL[col['format']]}"]
    if col['format'] == 'enum':
        lines.append('값: ' + ' / '.join(v['label'] for v in col['enumValues']))
    rv = rule_validation(columns, sheet, col)
    if rv is not None:
        lines.append('검사: ' + rv[3])
    rw = required_when_text(sheet, col)
    if rw:
        lines.append(f'조건: {rw} — 비면 칸이 빨갛게 됩니다')
    empty = empty_value_text(col, columns['boolLabels'])
    if empty is not None:
        lines.append(f'비우면: {empty}')
    return '\n'.join(lines)


def list_name(table: str) -> str:
    return f'L_{table}'


def range_ref(sheet_name: str, col: int, first: int, last: int) -> str:
    letter = get_column_letter(col)
    quoted = "'" + sheet_name.replace("'", "''") + "'"
    return f'{quoted}!${letter}${first}:${letter}${last}'


def add_list_validation(ws, cell_range: str, name: str, error: str, style: str) -> None:
    # OOXML 수식은 앞의 '=' 없이 저장한다
    dv = DataValidation(type='list', formula1=name, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = style
    dv.errorTitle = '값 확인'
    dv.error = error
    ws.add_data_validation(dv)
    dv.add(cell_range)


def add_rule_validation(ws, cell_range: str, rv: tuple[str, str, str | None, str]) -> None:
    kind, f1, f2, msg = rv
    if kind == 'custom':
        dv = DataValidation(type='custom', formula1=f1, allow_blank=True)
    else:
        dv = DataValidation(type=kind, operator='between', formula1=f1, formula2=f2, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorStyle = 'stop'
    dv.errorTitle = '값 확인'
    dv.error = msg
    ws.add_data_validation(dv)
    dv.add(cell_range)


def guide_doc(tool_root: Path) -> str | None:
    folder = tool_root / '계획_라벨링'
    found = sorted(p.name for p in folder.glob('지침서_v*.md') if '부록' not in p.name) if folder.is_dir() else []
    return f'계획_라벨링/{found[0]}' if found else None


def build_list_sheet(wb: Workbook, columns: dict) -> None:
    """숨긴 '목록' 시트: 열거 표마다 한 열 (1행 = 표 이름, 2행부터 표시 이름) + 참거짓 한 열"""
    ws = wb.create_sheet(LIST_SHEET)
    tables: dict[str, list[str]] = {}
    for s in columns['sheets']:
        for c in s['columns']:
            if c['format'] == 'enum' and c['enumTable'] not in tables:
                tables[c['enumTable']] = [v['label'] for v in c['enumValues']]
    tables_with_bool = [*tables.items(), ('BOOL', [columns['boolLabels']['true'], columns['boolLabels']['false']])]
    for i, (table, labels) in enumerate(tables_with_bool, start=1):
        ws.cell(row=1, column=i, value=table).font = Font(bold=True)
        for r, label in enumerate(labels, start=2):
            ws.cell(row=r, column=i, value=label)
        ws.column_dimensions[get_column_letter(i)].width = 30
        name = BOOL_LIST if table == 'BOOL' else list_name(table)
        wb.defined_names.add(DefinedName(name, attr_text=range_ref(LIST_SHEET, i, 2, len(labels) + 1)))
    ws.sheet_state = 'hidden'


def build_data_sheet(wb: Workbook, sheet: dict, columns: dict, rows: list[dict[int, object]]) -> None:
    ws = wb.create_sheet(sheet['name'])
    last = last_row(sheet)
    for i, col in enumerate(sheet['columns'], start=1):
        cell = ws.cell(row=1, column=i, value=col['header'])
        _, fill = header_status(col)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        comment = Comment(header_comment(columns, sheet, col), COMMENT_AUTHOR)
        comment.width = 380
        comment.height = 190
        cell.comment = comment
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = column_width(col)
        rng = data_range(sheet, i)
        fmt = col['format']
        if fmt == 'enum':
            add_list_validation(ws, rng, list_name(col['enumTable']), '목록에 있는 값만 고릅니다. 모르면 비웁니다.', 'stop')
        elif fmt == 'bool':
            add_list_validation(ws, rng, BOOL_LIST, "'예' 또는 '아니오'를 고릅니다.", 'stop')
        else:
            rv = rule_validation(columns, sheet, col)
            if rv is not None:
                add_rule_validation(ws, rng, rv)
        cf = required_when_formula(sheet, col)
        if cf is not None:
            ws.conditional_formatting.add(rng, FormulaRule(formula=[cf], fill=FILL_MISSING))
        # 글자 칸은 '@'(글자), 시각 칸은 날짜·시각 형식으로 미리 맞춘다 (엑셀이 '1' 을 숫자로, 번호를 날짜로 바꾸지 않게)
        if fmt in ('text', 'time'):
            number_format = '@' if fmt == 'text' else 'yyyy-mm-dd hh:mm'
            for r in range(FIRST_DATA_ROW, last + 1):
                ws.cell(row=r, column=i).number_format = number_format
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(sheet["columns"]))}{last}'
    for r, cells in enumerate(rows, start=FIRST_DATA_ROW):
        for c, value in cells.items():
            ws.cell(row=r, column=c, value=value)
    if sheet['id'] == 'physical':
        ws.sheet_properties.tabColor = PHYSICAL_TAB_COLOR
    if sheet['id'] == 'reader':
        ws.sheet_properties.tabColor = READER_TAB_COLOR


def add_reference_lists(wb: Workbook, columns: dict) -> None:
    done = set()
    for target_sheet, target_key, source_sheet, source_key, name in REFERENCE_LISTS:
        src = sheet_by_id(columns, source_sheet)
        if name not in done:
            wb.defined_names.add(
                DefinedName(name, attr_text=range_ref(src['name'], column_index(src, source_key), FIRST_DATA_ROW, LAST_DATA_ROW))
            )
            done.add(name)
        dst = sheet_by_id(columns, target_sheet)
        add_list_validation(
            wb[dst['name']],
            data_range(dst, column_index(dst, target_key)),
            name,
            f"'{src['name']}' 시트에 없는 값입니다. 그 시트에 먼저 적거나, 맞으면 그대로 둡니다.",
            'warning',
        )


def guide_title(blind: bool) -> str:
    return f'PAUT 결함 기록 양식 {FORM_VERSION} — ' + ('판독자용 (블라인드)' if blind else '대조용 (결함정보 포함)')


def build_guide_sheet(wb: Workbook, columns: dict, fingerprint: str, tool_root: Path, blind: bool) -> None:
    ws = wb.create_sheet(GUIDE_SHEET, 0)
    widths = {'A': 14, 'B': 30, 'C': 22, 'D': 18, 'E': 46, 'F': 80}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    row = 1

    def line(text: str = '', bold: bool = False, size: int | None = None, color: str | None = None) -> None:
        nonlocal row
        cell = ws.cell(row=row, column=1, value=text if text else None)
        if bold or size or color:
            cell.font = Font(bold=bold, size=size or 11, color=color)
        row += 1

    names = {s['id']: s['name'] for s in columns['sheets']}
    guide = guide_doc(tool_root)
    if guide is None:
        raise TemplateError("계획_라벨링/ 에 지침서(지침서_v*.md)가 없습니다. 지침서 경로 없이 양식을 만들지 않습니다.")
    line(guide_title(blind), bold=True, size=14)
    if not blind:
        line(
            "주의: 이 양식의 '물리결함' 시트에는 결함 위치·종류(설계값)가 들어 있습니다. 블라인드 판독자에게 주지 않습니다. "
            '판독자에게는 판독자용 양식(기록양식_v1.xlsx)을 줍니다.',
            bold=True,
            color='C00000',
        )
    line('OmniPC 로 .nde 를 판독해 지시를 적는 양식입니다 (라벨링 0단계). 칸은 앱의 기록 칸과 같아서, 앱이 이 파일을 그대로 가져와 앱에서 찍은 기록과 mm 로 비교합니다.')
    line(
        f"양식 {FORM_VERSION} · 라벨 스키마 v{columns['labelSchemaVersion']} · 칸 정의 지문 {fingerprint} "
        f"(칸 정의는 app/src/shared/labels/columns.ts 하나에서 나옵니다. 이 파일의 칸을 손으로 바꾸지 않습니다)"
    )
    line('대상: X3 12T-SV 4개 파일 (06_원시데이터/02_라벨완료/X3_12T-SV). 두께 12 mm 평판, 섹토리얼 Skew 90·270 두 그룹.')
    line()
    line('시트', bold=True)
    for s in columns['sheets']:
        if blind and s['id'] == 'physical':
            line(f"  {s['name']} — 판독자용 양식에서는 비어 있습니다 (대조 단계에서 씁니다).")
            continue
        line(f"  {s['name']} — {s['help']}")
    line(f'  {LIST_SHEET} (숨김) — 드롭다운 값. 고치지 않습니다.')
    line()
    line('쓰는 순서', bold=True)
    for text in [
        f"0. 파일을 복사해 내 이름의 파일로 저장합니다(예: 기록양식_v1_kim-l2_블라인드.xlsx). '{names['reader']}' 시트에 라벨러 ID 를 적고 판독 단계를 확인합니다. 파일 이름이 아니라 이 시트로 판독자를 가릅니다.",
        f"1. '{names['file']}' 시트를 확인합니다. 파일명·시편 ID·두께는 미리 채웠고, 기준점·방향은 가설(가설 정렬 = 예)입니다. 확인한 것만 고칩니다.",
        f"2. 파일을 열면 '{names['worktime']}' 시트에 판독 시작 시각을 적습니다.",
        f"3. C-scan 에서 지시를 찾고, 지시마다 '{names['indication']}' 시트에 한 줄을 적습니다. 같은 결함은 두 그룹에서 보여도 한 줄이고, '그룹' 칸에는 가장 잘 보인 그룹을 적습니다.",
        f"4. 결함이 아닌 루트·캡 형상 에코도 '{names['indication']}' 시트에 '기록 구분 = 형상 에코'로 적습니다. 한 줄 = 한 그룹 · 끊기지 않은 U 구간 하나(레그는 메모).",
        f"5. 다 본 구간은 '{names['review']}' 시트에 '검토 완료'로 적습니다(게이트 깊이는 펼친 깊이 그대로). 커플링 불량·누락 같은 구간은 '품질 불량', 보이지만 평가 기준에 못 미치는 응답은 '기준 미만 지시'로 적습니다.",
        f"6. 그 파일을 다 적으면 '{names['worktime']}' 시트에 판독 끝 시각, 쉰 시간, 라벨 추가 작업 시간(형상 에코·검토구간 적기, 게이트를 바꿔 다시 보기)을 적습니다.",
    ]:
        line('  ' + text)
    line()
    line('적는 규칙', bold=True)
    for text in [
        '칸에는 숫자만 적습니다. 단위는 머리글에 있습니다 (예: 스캔 시작 (mm) 칸에 12.5). 범위를 벗어나면 엑셀이 막습니다.',
        '스캔 위치는 OmniPC 스캔축 값 그대로 적습니다. 기준점으로 바꾸는 계산은 앱이 합니다.',
        '진폭은 소프트 이득 0 dB(파일에 저장된 이득 그대로)에서 읽습니다. 200 % 이상이면 포화 칸에 예.',
        "'Ref 대비 (dB, 앱 계산)' 칸은 비워 둡니다. OmniPC 의 AdBA 같은 읽기값은 Ref 기준이 아닐 수 있어 적지 않습니다(메모에는 적어도 됨).",
        '지시 시트의 깊이는 실제 깊이(표면에서 잰 값, 레그를 접은 값)이고 두께보다 클 수 없습니다. 결함 줄은 실제 깊이 위·아래를 꼭 적습니다.',
        '검토구간 시트의 게이트 깊이는 OmniPC 에 보이는 펼친 깊이 그대로입니다 (예: X3 게이트 A 7~28 → 레그 1~3).',
        '길이: 두께 8 mm 이상은 평가 기준(Ref −14 dB, X3 약 32 %FSH) 초과 구간, 8 mm 미만은 6 dB 강하 + 곡률 보정 (시방서 15절). 잰 그룹·각도·문턱도 적습니다.',
        "높이: '높이 입력 방식'을 위에서부터 고릅니다(상하 팁 → 팁 + 표면 → 노정 시간). −6 dB 높이는 시방서 밖 참고값입니다. 잴 수 없으면 '측정 불가'.",
        "목록 칸은 드롭다운에서 고릅니다. 모르면 비웁니다. 지시가 정말 있는지 자신 없으면 '존재 확신도'를 '불확실'로 고릅니다.",
        '칸이 빨갛게 되면 조건부 필수 칸이 빈 것입니다 (예: 기록 구분이 형상 에코인데 형상 에코 세부가 빔).',
        '날짜·시각은 2026-10-20 14:05 처럼 적습니다 (한국 시간).',
        f'드롭다운과 검사는 {DATA_ROWS}줄까지 걸어 두었습니다. 더 필요하면 위 줄을 복사해 붙입니다.',
    ]:
        line('  ' + text)
    line()
    line('색 표시 (머리글 칸 색)', bold=True)
    for text, fill in (
        ('  주황 = 필수 칸', FILL_REQUIRED),
        ('  파랑 = 비워도 됨 (앱이 가져올 때 채움)', FILL_APP),
        ('  회색 = 선택 칸 (조건부 필수 칸은 조건이 맞을 때 빈 칸이 빨갛게 됨)', FILL_OPTIONAL),
    ):
        ws.cell(row=row, column=1).fill = fill
        line(text)
    line('  머리글에 마우스를 올리면 설명(메모)이 보입니다.')
    line()
    if blind:
        line('블라인드 (설계값을 보지 않고 판독)', bold=True)
        line(f"  이 양식의 '{names['physical']}' 시트는 비어 있습니다. 결함정보 문서와 대조용 양식은 판독이 끝난 뒤에 받습니다.")
        line(f"  '{names['reader']}' 시트의 판독 단계 = 블라인드. 다 적으면 저장하고 더 고치지 않습니다(읽기 전용으로 둡니다).")
        line(f"  판독 전에 결함정보를 봤으면 '{names['indication']}' 시트의 '설계값 봤는지' 칸에 예를 적습니다.")
        line("  대조 단계: 이 파일을 복사해 …_대조.xlsx 를 만들고 판독 단계를 '설계값 대조'로 바꾼 뒤, 대조용 양식의 '물리결함' 시트를 봅니다(지침서 2.4절).")
    else:
        line('대조 단계 (설계값을 보고 고치거나 더함)', bold=True)
        line(f"  '{names['physical']}' 시트에 00_결함정보.docx 의 결함(종류·위치)을 미리 채웠습니다. 출처 '없음', 확인 수준 '미확인'입니다.")
        line(f"  '{names['reader']}' 시트의 판독 단계 = 설계값 대조. 이 단계의 줄은 앱이 모두 '설계값 봤음'으로 가져옵니다.")
        line("  대조에서 새로 더한 지시는 지시 번호를 C1, C2 … 로 적고 메모에 'docx 보고 추가'. 같은 결함은 '물리 결함 ID' 칸에 D1, D2 … 를 적습니다.")
        line('  결함정보.docx 의 위치가 결함 중심인지 시작인지, 기준점이 어디인지는 아직 모릅니다. 아시면 메모에 적어 주세요.')
    line()
    line('참고 문서 (10_데이터관리툴/ 기준)', bold=True)
    line(f'  지침서: {guide} (X3 예시 그림 부록은 판독이 끝난 뒤에 봅니다)')
    line('  계획: 계획_라벨링/계획.md · 규격: 계획_라벨링/규격_초안.md · 평가 정의: 계획_라벨링/평가_정의_v1.md · 신규 취득 규약: 계획_라벨링/취득_규약_v1.md')
    line('  시방서 요약: ../04_조사·보고/한난_PAUT시방서_핵심정리.md · 이 양식 다시 만들기: app/scripts/label-template/README.md')
    line()
    line('칸 목록', bold=True)
    header_row = row
    for i, h in enumerate(['시트', '머리글', '필수', '형식', '고를 수 있는 값 · 검사', '설명'], start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = Font(bold=True)
        c.border = HEADER_BORDER
    row += 1
    for s in columns['sheets']:
        for col in s['columns']:
            status, fill = header_status(col)
            values = ' / '.join(v['label'] for v in col['enumValues']) if col['format'] == 'enum' else ''
            if col['format'] == 'bool':
                values = f"{columns['boolLabels']['true']} / {columns['boolLabels']['false']}"
            rv = rule_validation(columns, s, col)
            if rv is not None:
                values = rv[3]
            req = '필수' if col['required'] else (required_when_text(s, col) or '')
            cells = [s['name'], col['header'], req, FORMAT_LABEL[col['format']], values, col['help']]
            for i, v in enumerate(cells, start=1):
                c = ws.cell(row=row, column=i, value=v)
                c.alignment = Alignment(wrap_text=i in (3, 5, 6), vertical='top')
            ws.cell(row=row, column=2).fill = fill
            row += 1
    ws.sheet_view.showGridLines = False


def build_workbook(columns: dict, prefill: dict, tool_root: Path, blind: bool) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    fingerprint = columns_fingerprint(columns)
    rows = mode_rows(columns, prefill, blind)
    build_list_sheet(wb, columns)
    for s in columns['sheets']:
        build_data_sheet(wb, s, columns, rows.get(s['id'], []))
    add_reference_lists(wb, columns)
    build_guide_sheet(wb, columns, fingerprint, tool_root, blind)
    # 숨긴 목록 시트는 맨 뒤로
    wb.move_sheet(LIST_SHEET, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(LIST_SHEET))
    wb.active = 0
    wb.properties.title = guide_title(blind)
    wb.properties.creator = 'PAUT 데이터 관리 · 라벨 양식 생성기'
    for name, value in [
        ('paut.labelForm', FORM_VERSION),
        ('paut.labelSchemaVersion', str(columns['labelSchemaVersion'])),
        ('paut.columnsFingerprint', fingerprint),
        ('paut.columnsSource', columns['source']),
        ('paut.blind', 'true' if blind else 'false'),
        ('paut.readingStage', STAGE_BLIND if blind else STAGE_RECONCILED),
    ]:
        wb.custom_doc_props.append(StringProperty(name=name, value=value))
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == GUIDE_SHEET
    return wb


# ───────────────────────────── 덮어쓰기 보호 ─────────────────────────────


def user_edits(path: Path, columns: dict, prefill: dict) -> list[str]:
    """이미 있는 양식에서 사람이 적은 것으로 보이는 칸 (미리 채운 값과 다른 칸). 몇 개만 돌려준다"""
    wb = load_workbook(path, read_only=True)
    allowed: dict[str, set[tuple[int, int, object]]] = {}
    for blind in (False, True):
        for sid, lines in mode_rows(columns, prefill, blind).items():
            sname = sheet_by_id(columns, sid)['name']
            for r, cells in enumerate(lines, start=FIRST_DATA_ROW):
                for c, v in cells.items():
                    allowed.setdefault(sname, set()).add((r, c, v))
    found: list[str] = []
    for s in columns['sheets']:
        if s['name'] not in wb.sheetnames:
            continue
        ws = wb[s['name']]
        for r, values in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW):
            for c, v in enumerate(values, start=1):
                if v is None or v == '':
                    continue
                if (r, c, v) not in allowed.get(s['name'], set()):
                    found.append(f"'{s['name']}' {get_column_letter(c)}{r} = {v!r}")
                    if len(found) >= 5:
                        return found
    return found


# ───────────────────────────── 다시 열어 확인 ─────────────────────────────


def verify(path: Path, columns: dict, prefill: dict, blind: bool) -> dict:
    wb = load_workbook(path)
    data_names = [s['name'] for s in columns['sheets']]
    expect_names = [GUIDE_SHEET, *data_names, LIST_SHEET]
    problems: list[str] = []
    if wb.sheetnames != expect_names:
        problems.append(f'시트 순서 {wb.sheetnames} ≠ {expect_names}')
    if wb[LIST_SHEET].sheet_state != 'hidden':
        problems.append("'목록' 시트가 숨김이 아닙니다.")
    defined = set(wb.defined_names.keys())
    rows = mode_rows(columns, prefill, blind)
    report = {'sheets': [], 'definedNames': len(defined), 'fingerprint': columns_fingerprint(columns)}
    for s in columns['sheets']:
        ws = wb[s['name']]
        last = last_row(s)
        headers = [ws.cell(row=1, column=i).value for i in range(1, len(s['columns']) + 1)]
        expected = [c['header'] for c in s['columns']]
        if headers != expected:
            problems.append(f"'{s['name']}' 머리글이 다릅니다: {headers}")
        if ws.cell(row=1, column=len(s['columns']) + 1).value is not None:
            problems.append(f"'{s['name']}' 머리글 뒤에 칸이 더 있습니다.")
        kinds = {'enum': 0, 'bool': 0, 'ref': 0, 'number': 0, 'time': 0}
        by_range: dict[str, list[tuple[str, str]]] = {}
        for dv in ws.data_validations.dataValidation:
            f1 = (dv.formula1 or '').lstrip('=')
            for rng in str(dv.sqref).split():
                by_range.setdefault(rng, []).append((dv.type, f1))
            if dv.type == 'list' and f1 not in defined:
                problems.append(f"'{s['name']}' 드롭다운 이름 {f1} 이 정의되지 않았습니다.")
        for i, c in enumerate(s['columns'], start=1):
            got = by_range.get(data_range(s, i), [])
            want = expected_validations(columns, s, c)
            if sorted(got) != sorted(want):
                problems.append(f"'{s['name']}' '{c['header']}' 유효성 {got} ≠ {want}")
            for t, f1 in want:
                if t == 'list' and f1 == BOOL_LIST:
                    kinds['bool'] += 1
                elif t == 'list' and f1 in {r[4] for r in REFERENCE_LISTS}:
                    kinds['ref'] += 1
                elif t == 'list':
                    kinds['enum'] += 1
                else:
                    kinds['time' if c['format'] == 'time' else 'number'] += 1
        # 조건부 필수 칸의 조건부 서식
        cf_got = {}
        for cf_range in ws.conditional_formatting:
            for rule in cf_range.rules:
                cf_got[str(cf_range.sqref)] = list(rule.formula)
        cf_want = {}
        for i, c in enumerate(s['columns'], start=1):
            f = required_when_formula(s, c)
            if f is not None:
                cf_want[data_range(s, i)] = [f]
        if cf_got != cf_want:
            problems.append(f"'{s['name']}' 조건부 서식 {cf_got} ≠ {cf_want}")
        comments = sum(1 for i in range(1, len(s['columns']) + 1) if ws.cell(row=1, column=i).comment is not None)
        if comments != len(s['columns']):
            problems.append(f"'{s['name']}' 머리글 메모 {comments}/{len(s['columns'])}")
        if ws.freeze_panes != 'B2':
            problems.append(f"'{s['name']}' 고정 위치 {ws.freeze_panes}")
        required = [c['header'] for c in s['columns'] if c['required']]
        for i, c in enumerate(s['columns'], start=1):
            color = ws.cell(row=1, column=i).fill.fgColor.rgb
            if c['required'] and not str(color).endswith(FILL_REQUIRED.fgColor.rgb[-6:]):
                problems.append(f"'{s['name']}' 필수 칸 '{c['header']}' 색이 다릅니다.")
        filled = rows.get(s['id'], [])
        for r, cells in enumerate(filled, start=FIRST_DATA_ROW):
            for col, value in cells.items():
                got = ws.cell(row=r, column=col).value
                if got != value:
                    problems.append(f"'{s['name']}' {r}행 {col}열: {got!r} ≠ {value!r}")
        if ws.cell(row=FIRST_DATA_ROW + len(filled), column=1).value is not None:
            problems.append(f"'{s['name']}' 미리 채운 줄 뒤에 값이 더 있습니다.")
        if blind and s['id'] == 'physical' and ws.max_row > 1:
            if any(v is not None for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True) for v in row):
                problems.append("판독자용 양식의 '물리결함' 시트에 값이 있습니다.")
        report['sheets'].append(
            {
                'name': s['name'],
                'columns': len(s['columns']),
                'required': len(required),
                'dropdowns': kinds['enum'] + kinds['bool'] + kinds['ref'],
                'dropdownsByKind': {k: kinds[k] for k in ('enum', 'bool', 'ref')},
                'numberChecks': kinds['number'],
                'timeChecks': kinds['time'],
                'conditional': len(cf_want),
                'comments': comments,
                'prefilledRows': len(filled),
                'rows': last - FIRST_DATA_ROW + 1,
            }
        )
    # 목록 시트 값 = 칸 정의의 표시 이름
    lists = wb[LIST_SHEET]
    for i in range(1, lists.max_column + 1):
        table = lists.cell(row=1, column=i).value
        labels = [lists.cell(row=r, column=i).value for r in range(2, lists.max_row + 1)]
        labels = [v for v in labels if v is not None]
        if table == 'BOOL':
            want_labels = [columns['boolLabels']['true'], columns['boolLabels']['false']]
        else:
            col = next(c for s in columns['sheets'] for c in s['columns'] if c.get('enumTable') == table)
            want_labels = [v['label'] for v in col['enumValues']]
        if labels != want_labels:
            problems.append(f"'목록' {table}: {labels} ≠ {want_labels}")
    report['listTables'] = lists.max_column
    # 안내 시트 문구: 오래된 '작성 중' 문구가 없고, 판독자용에는 결함정보 내용을 말하지 않는다
    guide_text = '\n'.join(str(v) for row in wb[GUIDE_SHEET].iter_rows(values_only=True) for v in row if v is not None)
    if STALE_GUIDE_TEXT in guide_text:
        problems.append(f"안내 시트에 '{STALE_GUIDE_TEXT}' 문구가 남아 있습니다.")
    if blind and ('미리 채웠습니다. 출처' in guide_text or '결함정보.docx 의 결함' in guide_text):
        problems.append('판독자용 안내 시트에 결함정보를 채웠다는 문구가 있습니다.')
    if not blind and '블라인드 판독자에게 주지 않습니다' not in guide_text:
        problems.append('대조용 안내 시트에 블라인드 경고가 없습니다.')
    props = {p.name: p.value for p in wb.custom_doc_props}
    if props.get('paut.columnsFingerprint') != report['fingerprint']:
        problems.append(f"문서 속성 지문 {props.get('paut.columnsFingerprint')} ≠ {report['fingerprint']}")
    if props.get('paut.blind') != ('true' if blind else 'false'):
        problems.append(f"문서 속성 paut.blind = {props.get('paut.blind')}")
    if problems:
        raise TemplateError('만든 양식을 다시 열어 보니 다릅니다:\n  ' + '\n  '.join(problems))
    return report


# ───────────────────────────── 실행 ─────────────────────────────


def main(argv: list[str] | None = None) -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description='PAUT 결함 기록 양식(엑셀) 생성')
    ap.add_argument('--columns', required=True, type=Path, help='칸 정의 JSON (templateColumns.test.ts 가 씀)')
    ap.add_argument('--out', required=True, type=Path, help='만들 xlsx 경로')
    ap.add_argument('--prefill', type=Path, default=here / 'prefill_x3.json', help='미리 채울 값 JSON')
    ap.add_argument('--docx', type=Path, default=None, help='(선택) 00_결함정보.docx — prefill 물리결함과 대조')
    ap.add_argument('--tool-root', type=Path, default=here.parents[2], help='10_데이터관리툴 폴더 (지침서 찾기용)')
    ap.add_argument('--blind', action='store_true', help="판독자용: '물리결함' 시트를 비우고 판독 단계 = 블라인드")
    ap.add_argument('--force', action='store_true', help='사람이 적은 값이 있어도 --out 을 덮어씀')
    args = ap.parse_args(argv)
    tmp = args.out.with_name(f'.{args.out.stem}.만드는중.xlsx')
    try:
        columns = load_columns(args.columns)
        prefill = json.loads(args.prefill.read_text(encoding='utf-8'))
        if args.docx is not None:
            n = check_docx(args.docx, prefill)
            print(f'결함정보.docx 대조: {n}건 같음')
        if args.out.exists() and not args.force:
            edits = user_edits(args.out, columns, prefill)
            if edits:
                raise TemplateError(
                    f'{args.out} 에 사람이 적은 것으로 보이는 값이 있어 덮어쓰지 않습니다 (예: {", ".join(edits)}). '
                    '다른 --out 을 주거나, 덮어써도 되면 --force 를 더합니다.'
                )
        wb = build_workbook(columns, prefill, args.tool_root, args.blind)
        args.out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(tmp)
        report = verify(tmp, columns, prefill, args.blind)
        os.replace(tmp, args.out)
    except TemplateError as e:
        if tmp.exists():
            tmp.unlink()  # 이번 실행이 만든 임시 파일만 지운다 (--out 은 그대로)
        print(f'오류: {e}', file=sys.stderr)
        return 1
    print(f"저장: {args.out} ({'판독자용·블라인드' if args.blind else '대조용·결함정보 포함'})")
    print(f"칸 정의 지문 {report['fingerprint']} · 정의된 이름 {report['definedNames']}개 · 목록 표 {report['listTables']}개")
    for s in report['sheets']:
        k = s['dropdownsByKind']
        print(
            f"  {s['name']}: 칸 {s['columns']} (필수 {s['required']}) · 줄 {s['rows']} · 드롭다운 {s['dropdowns']} "
            f"(열거 {k['enum']}, 예/아니오 {k['bool']}, 다른 시트 참조 {k['ref']}) · 숫자 검사 {s['numberChecks']} · "
            f"시각 검사 {s['timeChecks']} · 조건부 필수 {s['conditional']} · 머리글 메모 {s['comments']} · 미리 채운 줄 {s['prefilledRows']}"
        )
    return 0


if __name__ == '__main__':
    sys.exit(main())
